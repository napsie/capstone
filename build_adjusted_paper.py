from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\PLPASIG\.codex\attachments\985e763b-2540-4368-bae8-b77d47df313a\pasted-text.txt")
OUT = Path(r"D:\xampp1\htdocs\capstone\SENIORLINK_Adjusted_Paper.docx")

raw = SRC.read_text(encoding="utf-8")
for bad, good in {
    "Â§": "•", "â—": "•", "â€“": "–", "â€™": "’", "â€œ": "“", "â€": "”",
}.items():
    raw = raw.replace(bad, good)
lines = [x.strip() for x in raw.splitlines()]

def replace_range(start, end, replacement):
    global lines
    a = next(i for i, x in enumerate(lines) if x == start)
    b = next(i for i, x in enumerate(lines[a + 1:], a + 1) if x == end)
    lines[a:b] = [start] + replacement

rrl = [
    "Digital Profiling",
    "Digital profiling refers to the systematic collection, organization, updating, and retrieval of an individual’s demographic and service-related information through an electronic platform. In public administration, a centralized profile can reduce duplicate records, support faster retrieval, and allow authorized offices to work from consistent information. For senior citizen services, a digital profile may include identity details, residency information, contact information, service history, and submitted documentary requirements. The value of profiling, however, depends on clear data standards, controlled access, regular updating procedures, and compliance with the Data Privacy Act of 2012. SENIORLINK applies digital profiling as an administrative procedure in which records are encoded once, checked against existing entries, validated by authorized personnel, and maintained in a centralized repository.",
    "Senior Citizen Services",
    "Senior citizen services include identification, social pension assistance, burial assistance, cash benefits, referrals, and other welfare programs administered by local and national government offices. These services require accurate identification, complete supporting documents, consistent eligibility assessment, and timely coordination between barangay help desks and the main Office of Senior Citizens Affairs. Paper-dependent procedures can create repeated visits, delayed verification, and difficulty locating application records. A digital service platform should therefore support the existing legal and administrative process while improving accessibility for older persons, including bedridden and low-mobility applicants who may need an authorized representative.",
    "Rule-Based Systems",
    "A rule-based system applies explicit if–then conditions to data supplied by users or retrieved from authorized records. Unlike probabilistic decision models, its results can be traced to a defined requirement, ordinance, or office procedure. In the proposed system, rules are used for preliminary checking of age, residency, documentary completeness, filing periods, and other approved eligibility conditions. A rule match does not replace the authority of OSCA personnel. Instead, the system records the result, identifies missing or inconsistent information, and routes exceptional cases for human review. This approach supports consistent screening while preserving administrative accountability.",
    "Workflow Automation",
    "Workflow automation coordinates tasks, records, notifications, and responsible personnel according to an approved sequence of work. In government service delivery, automation is useful when an application must pass through multiple offices and checks before approval. SENIORLINK uses rule-based workflow routing to move an application through application submission, requirement checking, verification, document validation, authorized approval, and release. An incomplete or inconsistent application is returned to the appropriate step with a recorded reason. Each routing action is time-stamped and associated with the responsible user to improve monitoring and accountability.",
    "Business Process Management",
    "Business Process Management is the disciplined analysis, design, implementation, monitoring, and improvement of organizational processes. Its purpose is not simply to computerize an existing activity but to establish a clear and repeatable procedure. Applied to senior citizen services, it requires the researchers to document the current process, identify delays and duplicated work, define responsible personnel, set decision rules, and measure whether the redesigned process improves service delivery. The proposed system therefore treats technology as an enabler of a standardized OSCA process rather than as the solution by itself.",
    "Electronic Document Management",
    "Electronic document management covers the controlled capture, classification, storage, retrieval, versioning, retention, and disposal of digital documents. For OSCA transactions, this includes scanned application forms, proof of identity, residency documents, authorization records, and other service-specific requirements. Documents should be linked to the correct applicant and transaction, protected through role-based access, and supported by an audit history. Electronic copies do not automatically replace originals or legally required hard copies; retention and disposal must follow applicable government and Commission on Audit requirements. Within SENIORLINK, document management supports requirement checking, validation, search, status monitoring, and authorized retrieval.",
    "Integration of the Reviewed Concepts",
    "The reviewed concepts show that effective senior citizen service delivery requires more than a centralized database. Digital profiling supplies organized citizen information; senior citizen service standards define the administrative context; rule-based logic supports consistent preliminary checks; workflow automation routes work among responsible offices; Business Process Management standardizes and improves the procedure; and electronic document management protects and retrieves supporting records. The research gap addressed by SENIORLINK is the lack of a locally focused platform that connects these elements across barangay help desks and the main OSCA office while retaining human authorization for final administrative decisions.",
]

replace_range("Computational Modeling and Logic-Driven Algorithms", "Advanced Analytics and Ethical Considerations", rrl)

replace_range("Algorithmic Logic and Ethical Implementation", "Comprehensive Synthesis and Research Gap", [
    "Rule-Based Logic and Ethical Implementation",
    "The literature indicates that a contemporary social welfare information system must combine standardized procedures, transparent decision rules, useful records, and strong ethical safeguards. Rule-based logic is appropriate for preliminary compliance checking because each result can be traced to an approved condition. The rules must be documented, version-controlled, tested against valid and invalid cases, and reviewed whenever an ordinance or office procedure changes.",
    "SENIORLINK uses rule-based logic for eligibility screening and workflow routing. The system records which rule was applied, the data used, the result, the responsible user, and any manual action taken. Final approval remains with authorized OSCA personnel. This arrangement supports consistency and auditability without allowing software to make an unreviewable final administrative decision.",
])

replace_range("Statement of the Problem", "Scope and Limitation of the Study", [
    "The Office of Senior Citizens Affairs (OSCA) and barangay senior citizen help desks manage citizen profiles, documentary requirements, benefit applications, and service records. Existing paper-based and decentralized procedures can cause delayed processing, duplicate or inconsistent records, difficulty locating documents, uneven application of requirements, and barriers for senior citizens with limited mobility.",
    "The study addresses these problems by designing both a standardized administrative procedure and a supporting information system. The proposed procedure defines how applications are submitted, checked for completeness, verified against existing records, validated by authorized personnel, routed for correction or approval, and recorded for monitoring. Technology supports these activities but does not replace official judgment or required controls.",
    "Specifically, the study seeks to answer the following questions:",
    "1. How can a standardized digital profiling procedure improve the organization, accessibility, accuracy, security, and updating of senior citizen information across barangay help desks and the main OSCA office?",
    "2. How can documented rule-based checking procedures improve the consistency and efficiency of evaluating eligibility, document completeness, and compliance with approved local government requirements?",
    "3. How can rule-based workflow routing standardize application submission, requirement checking, verification, document validation, correction, approval, and release?",
    "4. How can an electronic document management procedure improve document capture, classification, retrieval, access control, retention, and auditability?",
    "5. How can representative authorization, role-based access, and application tracking improve secure access to services for bedridden and low-mobility senior citizens?",
    "6. How acceptable is the proposed process and system in terms of functionality, usability, reliability, security, efficiency, and support for established OSCA procedures?",
    "SOP-to-Solution Alignment",
    "The first problem is addressed through a profiling standard covering required fields, duplicate checking, controlled updating, and access permissions. The second is addressed through documented and testable rules derived from approved requirements. The third is addressed through a defined routing procedure with responsible personnel and return paths. The fourth is addressed through document classification, validation, retention, and retrieval controls. The fifth is addressed through verified representative authorization and traceable access. The sixth is addressed through user evaluation using criteria aligned with the study objectives.",
])

replace_range("Algorithm", "Software Engineering (SE) Paradigm", [
    "Rule-Based Workflow Routing",
    "The architecture connects the barangay help desk, the central application server, the rule repository, the electronic document repository, the senior citizen profile database, and the OSCA administrative dashboard. The barangay help desk captures applications and supporting records. The central server applies approved checking and routing rules, while the databases store profiles, documents, transactions, and audit logs. The OSCA dashboard allows authorized personnel to review exceptions, validate records, approve eligible applications, and monitor pending work.",
    "Application Submission",
    "The applicant or verified representative submits the required information and documents through the designated barangay help desk. The system creates a transaction reference, records the submission date and receiving officer, links the submission to an existing or newly created profile, and routes the application to requirement checking.",
    "Requirement Checking",
    "The system compares the submitted fields and documents with the approved checklist for the selected service. Complete submissions proceed to verification. Incomplete submissions are returned for completion with the missing requirements, responsible office, date, and remarks recorded. This routing rule prevents incomplete applications from entering the formal verification queue.",
    "Verification",
    "Authorized personnel compare the applicant’s information with existing profiles and transaction records. The procedure checks for possible duplicates, inconsistent identity information, residency discrepancies, prior applications, and other approved eligibility conditions. A clear match proceeds to document validation. A mismatch or uncertain case is routed to manual review; the system does not make the final determination.",
    "Document Validation",
    "Authorized personnel inspect submitted documents for readability, consistency, validity, and compliance with service-specific requirements. Valid documents proceed to approval. Invalid, expired, unreadable, or inconsistent documents are returned for correction with a documented reason. The system preserves the validation result and audit information for each document.",
    "Approval, Release, and Recording",
    "An authorized OSCA official reviews applications that passed the preceding checks and records the final decision. Approved applications are prepared for the appropriate service or benefit, while denied applications retain the reason and supporting review record. The system updates the centralized profile, records the responsible official and time of action, and makes the current status available to authorized users.",
    "Routing Rules and Controls",
    "Routing is determined by explicit conditions: if requirements are incomplete, return for completion; if identity data are inconsistent, route for manual verification; if a document is invalid, return for replacement; if all checks are satisfied, route to authorized approval. Every routing action is logged. Rules must be based on verified ordinances and office procedures, approved by responsible officials, tested before deployment, and updated through controlled change management.",
])

joined = "\n".join(lines)
joined = re.sub(r"\bFinite State Machine(?: \(FSM\))?\b", "rule-based workflow routing", joined, flags=re.I)
joined = re.sub(r"\bfinite state machines\b", "rule-based workflow routing", joined, flags=re.I)
joined = re.sub(r"\bFSM(?:-based)?\b", "rule-based", joined)
joined = joined.replace("Using Rule-Based Logic and rule-based workflow routing Workflow Automation", "Using Rule-Based Workflow Automation")
joined = joined.replace("Rule-Based Logic and rule-based workflow routing", "Rule-Based Workflow Automation")
joined = joined.replace("rule-based workflow routing workflow", "rule-based workflow")
lines = [x.strip() for x in joined.splitlines()]

# Remove obsolete mathematical/state-machine remnants and empty clutter.
lines = [x for x in lines if not any(k in x for k in [
    "formal quintuple", "M = (", "finite set of operational application states",
    "set of administrative input triggers", "initial intake state:", "final accepting state",
    "state transition function (Q",
])]

headings = {
    "CHAPTER 1": 0, "INTRODUCTION": 0, "CHAPTER 2": 0, "METHODOLOGY": 0,
    "The Problem and its Background": 1, "Review of Related Literature": 1,
    "Significance of the Study": 1, "Statement of the Problem": 1,
    "Scope and Limitation of the Study": 1, "Scope of the Study": 2,
    "Limitations of the Study": 2, "Introduction": 1, "2.2 Research Design": 1,
    "Architectural Framework": 1, "Rule-Based Workflow Routing": 1,
    "Software Engineering (SE) Paradigm": 1, "Population and Sampling": 1,
    "Population": 2, "Sampling Method": 2, "Sample Size": 2,
    "Data Collection Methods": 1, "Instruments": 2, "Procedure": 2,
    "Ethical Considerations": 1, "Summary": 1,
}
for h in rrl[::2] + ["Integration of the Reviewed Concepts", "Rule-Based Logic and Ethical Implementation",
                     "SOP-to-Solution Alignment", "Application Submission", "Requirement Checking",
                     "Verification", "Document Validation", "Approval, Release, and Recording",
                     "Routing Rules and Controls"]:
    headings[h] = 2

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name, normal.font.size = "Times New Roman", Pt(12)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.first_line_indent = Inches(.5)
for n, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)]:
    s = styles[n]
    s.font.name, s.font.size, s.font.bold = "Times New Roman", Pt(size), True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(12), Pt(6)
    s.paragraph_format.keep_with_next = True

for line in lines:
    if not line:
        continue
    if line in ("CHAPTER 1", "CHAPTER 2"):
        if len(doc.paragraphs) > 0 and line == "CHAPTER 2":
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold, r.font.name, r.font.size = True, "Times New Roman", Pt(14)
        continue
    if line in ("INTRODUCTION", "METHODOLOGY"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold, r.font.name, r.font.size = True, "Times New Roman", Pt(14)
        continue
    if line in headings:
        p = doc.add_heading(line, level=headings[line])
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if headings[line] else WD_ALIGN_PARAGRAPH.CENTER
        continue
    bullet = line.startswith("•")
    numbered = bool(re.match(r"^\d+\.\s", line))
    if bullet:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(line.lstrip("•").strip())
    elif numbered:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Inches(.25)
        p.paragraph_format.first_line_indent = Inches(-.25)
    else:
        p = doc.add_paragraph(line)
    p.paragraph_format.widow_control = True

# Footer page field.
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)

doc.core_properties.title = "SENIORLINK Adjusted Paper"
doc.core_properties.subject = "Revised capstone paper with rule-based workflow automation"
doc.core_properties.author = "Research Team"
doc.save(OUT)
print(OUT)
